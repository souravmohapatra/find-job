import io
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, date, datetime, timedelta
from typing import Any
from uuid import UUID, uuid4

from app.config import get_settings
from app.db import get_engine
from app.domain.profile import Employment, Facts, experience
from app.profile_models import Document, Operation, ProfileDraft, ProfileRevision
from app.services.operations import process_operation
from app.tasks import worker
from docx import Document as Docx
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

CV = (
    b"Synthetic product designer\nAcme example studio, 2020-2024\n"
    b"Designed accessible mobile applications.\nexample@example.invalid\n"
)


def auth(client: TestClient) -> dict[str, str]:
    response = client.post(
        "/api/v1/auth/login",
        headers={"Origin": "https://testserver"},
        json={"username": "owner", "password": "synthetic-test-password"},
    )
    assert response.status_code == 200
    return {"Origin": "https://testserver", "X-CSRF-Token": response.json()["csrf_token"]}


def upload(
    client: TestClient, headers: dict[str, str], content: bytes = CV, key: str | None = None
) -> dict[str, str]:
    response = client.post(
        "/api/v1/profile/documents",
        headers={**headers, "Idempotency-Key": key or str(uuid4())},
        files={"file": ("private-name.txt", content, "application/octet-stream")},
    )
    assert response.status_code == 202, response.text
    return response.json()


def draft_for(client: TestClient, result: dict[str, str]) -> dict[str, Any]:
    process_operation(UUID(result["operation_id"]))
    assert client.get(result["status_url"]).json()["state"] == "succeeded"
    draft = client.get("/api/v1/profile").json()["drafts"][0]
    return client.get("/api/v1/profile/drafts/" + draft["id"]).json()


def reviewed(draft: dict[str, Any]) -> dict[str, Any]:
    return {
        "facts": {**draft["facts"], "summary": "Reviewed synthetic designer experience"},
        "preferences": {
            **draft["preferences"],
            "desired_roles": ["Product design"],
            "domain_mode": "open",
        },
    }


def test_cv_review_activation_and_immutable_concurrent_edits(client: TestClient) -> None:
    headers = auth(client)
    result = upload(client, headers)
    assert client.get("/api/v1/profile").json()["revision"] == 0
    draft = draft_for(client, result)
    assert draft["preferences"]["domain_mode"] == "open"
    assert draft["facts"]["summary"] == "Synthetic product designer"
    assert all("example@example.invalid" not in span["text"] for span in draft["spans"])
    body = reviewed(draft)
    patch = client.patch(
        "/api/v1/profile/drafts/" + draft["id"], headers={**headers, "If-Match": '"1"'}, json=body
    )
    assert patch.status_code == 200
    activated = client.post(
        "/api/v1/profile/drafts/" + draft["id"] + "/activate",
        headers={**headers, "If-Match": '"0"'},
        json={"draft_version": 2, "reviewed": True},
    )
    assert activated.status_code == 200 and activated.json()["revision"] == 1
    assert (
        client.patch(
            "/api/v1/profile", headers={**headers, "If-Match": '"0"'}, json=body
        ).status_code
        == 409
    )
    body["preferences"]["domain_mode"] = "expand"
    body["preferences"]["domains"] = ["Education"]
    assert (
        client.patch("/api/v1/profile", headers={**headers, "If-Match": '"1"'}, json=body).json()[
            "revision"
        ]
        == 2
    )
    with Session(get_engine()) as db:
        versions = list(db.scalars(select(ProfileRevision).order_by(ProfileRevision.revision)))
        assert len(versions) == 2 and versions[0].preferences["domain_mode"] == "open"
        assert versions[1].preferences["domain_mode"] == "expand"
        assert (
            db.scalar(
                select(func.count())
                .select_from(Operation)
                .where(Operation.kind == "profile_reassess")
            )
            == 2
        )


def test_idempotency_and_duplicate_workers(client: TestClient) -> None:
    headers = auth(client)
    key = str(uuid4())
    first = upload(client, headers, key=key)
    second = upload(client, headers, key=key)
    assert first == second
    conflict = client.post(
        "/api/v1/profile/documents",
        headers={**headers, "Idempotency-Key": key},
        files={"file": ("cv.txt", b"Different content")},
    )
    assert conflict.status_code == 409
    with ThreadPoolExecutor(max_workers=2) as executor:
        list(executor.map(process_operation, [UUID(first["operation_id"])] * 2))
    with Session(get_engine()) as db:
        assert db.scalar(select(func.count()).select_from(Document)) == 1
        assert db.scalar(select(func.count()).select_from(ProfileDraft)) == 1
        operation = db.get(Operation, UUID(first["operation_id"]))
        assert operation and operation.state == "succeeded" and operation.attempts == 1


def test_replacement_keeps_preferences_and_conflicts_with_newer_profile(client: TestClient) -> None:
    headers = auth(client)
    body = {
        "facts": {"summary": "Reviewed nurse experience"},
        "preferences": {
            "desired_roles": ["Nursing"],
            "domain_mode": "expand",
            "domains": ["Community health"],
        },
    }
    assert (
        client.patch(
            "/api/v1/profile", headers={**headers, "If-Match": '"0"'}, json=body
        ).status_code
        == 200
    )
    result = upload(
        client, headers, b"Synthetic registered nurse\nClinical care and patient education."
    )
    draft = draft_for(client, result)
    assert draft["base_revision"] == 1 and draft["preferences"]["desired_roles"] == ["Nursing"]
    assert draft["base_profile"]["facts"]["summary"] == "Reviewed nurse experience"
    assert (
        client.patch(
            "/api/v1/profile", headers={**headers, "If-Match": '"1"'}, json=body
        ).status_code
        == 200
    )
    response = client.post(
        "/api/v1/profile/drafts/" + draft["id"] + "/activate",
        headers={**headers, "If-Match": '"2"'},
        json={"draft_version": 1, "reviewed": True},
    )
    assert response.status_code == 409
    assert client.get("/api/v1/profile").json()["revision"] == 2


def test_required_review_and_evidence_validation(client: TestClient) -> None:
    headers = auth(client)
    draft = draft_for(client, upload(client, headers))
    url = "/api/v1/profile/drafts/" + draft["id"]
    body = reviewed(draft)
    body["facts"]["skills"] = [
        {
            "value": "Invented skill",
            "origin": "extracted",
            "evidence": [{"span_id": "s999", "quote": "missing"}],
        }
    ]
    assert client.patch(url, headers={**headers, "If-Match": '"1"'}, json=body).status_code == 200
    assert (
        client.post(
            url + "/activate",
            headers={**headers, "If-Match": '"0"'},
            json={"draft_version": 2, "reviewed": True},
        ).status_code
        == 422
    )
    body["facts"]["skills"] = []
    body["preferences"]["desired_roles"] = []
    assert (
        client.patch(
            "/api/v1/profile", headers={**headers, "If-Match": '"0"'}, json=body
        ).status_code
        == 422
    )
    body["preferences"]["desired_roles"] = ["Designer"]
    body["preferences"]["domain_mode"] = "focus"
    body["preferences"]["domains"] = []
    assert (
        client.patch(
            "/api/v1/profile", headers={**headers, "If-Match": '"0"'}, json=body
        ).status_code
        == 422
    )


def test_document_limits_invalid_pdf_and_docx_text(client: TestClient) -> None:
    headers = auth(client)
    assert (
        client.post(
            "/api/v1/profile/documents",
            headers=headers,
            files={"file": ("huge.txt", b"x" * (10 * 1024 * 1024 + 1))},
        ).status_code
        == 413
    )
    assert (
        client.post(
            "/api/v1/profile/documents", headers=headers, files={"file": ("bad.txt", b"\x00\x01")}
        ).status_code
        == 422
    )
    pdf = PdfWriter()
    pdf.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    pdf.write(buffer)
    result = upload(client, headers, buffer.getvalue())
    process_operation(UUID(result["operation_id"]))
    failure = client.get(result["status_url"]).json()
    assert failure["state"] == "failed" and "OCR" in failure["error_message"]
    doc = Docx()
    doc.add_paragraph("Synthetic mechanical engineer")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Designed a cooling system"
    buffer = io.BytesIO()
    doc.save(buffer)
    draft = draft_for(client, upload(client, headers, buffer.getvalue()))
    assert any("cooling system" in span["text"] for span in draft["spans"])


def test_private_download_delete_and_profile_purge(client: TestClient) -> None:
    headers = auth(client)
    result = upload(client, headers)
    draft_for(client, result)
    url = "/api/v1/profile/documents/" + result["resource_id"]
    assert client.get(url + "/download").content == CV
    assert client.delete(url, headers=headers).status_code == 204
    assert client.get(url + "/download").status_code == 404
    with Session(get_engine()) as db:
        document = db.get(Document, UUID(result["resource_id"]))
        assert document and document.spans is None
    assert (
        client.post(
            "/api/v1/profile/purge", headers=headers, json={"confirmation": "no"}
        ).status_code
        == 422
    )
    pending = upload(client, headers)
    assert (
        client.post(
            "/api/v1/profile/purge", headers=headers, json={"confirmation": "DELETE PROFILE"}
        ).status_code
        == 204
    )
    process_operation(UUID(pending["operation_id"]))
    assert client.get(pending["status_url"]).json()["state"] == "cancelled"
    assert client.get("/api/v1/profile").json()["drafts"] == []
    assert list(get_settings().document_dir.iterdir()) == []


def test_real_queue_dispatch_and_restart_recovery(client: TestClient) -> None:
    headers = auth(client)
    result = upload(client, headers)
    with Session(get_engine()) as db, db.begin():
        op = db.get(Operation, UUID(result["operation_id"]))
        assert op
        op.state = "running"
        op.lease_until = datetime.now(UTC) - timedelta(minutes=1)
        op.lease_token = uuid4()
        op.attempts = 1
    with worker.app.open():
        assert worker.dispatch() == 1
        assert worker.dispatch() == 0
    worker.app.run_worker(
        queues=["ai"], wait=False, install_signal_handlers=False, delete_jobs="successful"
    )
    state = client.get(result["status_url"]).json()
    assert state["state"] == "succeeded" and state["attempts"] == 2


def test_experience_union_uncertainty_and_override() -> None:
    roles = [
        Employment(employer="A", title="Engineer", start="2020-01-01", end="2022-12-31"),
        Employment(employer="B", title="Engineer", start="2021-01-01", end="2023-12-31"),
    ]
    value = experience(Facts(employment=roles), date(2024, 1, 1))
    assert value["minimum_years"] == 4.0 and value["maximum_years"] == 4.0
    roles.append(Employment(employer="C", title="Engineer"))
    assert experience(Facts(employment=roles))["maximum_years"] is None
    assert experience(Facts(experience_years_override=7))["basis"] == "user_override"
    uncertain = experience(
        Facts(employment=[Employment(employer="A", title="Nurse", start="2020", end="2021")]),
        date(2024, 1, 1),
    )
    assert uncertain["minimum_years"] == 0 and uncertain["maximum_years"] == 2.1


def test_multi_page_pdf_and_retention_preserves_pending_drafts(client: TestClient) -> None:
    from app.services.retention import expire_approved_documents
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for heading in ("Synthetic researcher", "Employment and qualifications"):
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 40 740 Td ({heading}) Tj ET".encode())
        page[NameObject("/Contents")] = stream
    buffer = io.BytesIO()
    writer.write(buffer)
    headers = auth(client)
    approved = upload(client, headers, buffer.getvalue())
    draft = draft_for(client, approved)
    assert [span["page"] for span in draft["spans"]] == [1, 2]
    pending = upload(client, headers)
    retained = upload(client, headers, b"Synthetic retained CV")
    with Session(get_engine()) as db, db.begin():
        for item in (approved, pending, retained):
            document = db.get(Document, UUID(item["resource_id"]))
            assert document
            document.approved_at = datetime.now(UTC) - timedelta(days=31)
        for item in (approved, retained):
            row = db.scalar(
                select(ProfileDraft).where(ProfileDraft.document_id == UUID(item["resource_id"]))
            )
            assert row
            row.status = "activated"
        document = db.get(Document, UUID(retained["resource_id"]))
        assert document
        document.retain_original = True
    assert expire_approved_documents() == 1
    with Session(get_engine()) as db:
        original = db.get(Document, UUID(approved["resource_id"]))
        assert original
        assert original.spans is None and original.deleted_at is not None
        assert db.get(Document, UUID(pending["resource_id"])).deleted_at is None
        assert db.get(Document, UUID(retained["resource_id"])).deleted_at is None
    assert (
        client.get("/api/v1/profile/documents/" + approved["resource_id"] + "/download").status_code
        == 404
    )


def test_stalled_queue_job_is_recovered_without_duplicate_document(client: TestClient) -> None:
    import asyncio

    result = upload(client, auth(client))
    with worker.app.open():
        worker.dispatch()

    async def simulate_worker_loss() -> None:
        async with worker.app.open_async():
            identifier = await worker.app.job_manager.register_worker()
            job = await worker.app.job_manager.fetch_job(queues=["ai"], worker_id=identifier)
            assert job is not None
            await worker.app.job_manager.unregister_worker(identifier)
            assert await worker.recover_stalled_jobs() == 1

    asyncio.run(simulate_worker_loss())
    worker.app.run_worker(
        queues=["ai"], wait=False, install_signal_handlers=False, delete_jobs="successful"
    )
    assert client.get(result["status_url"]).json()["state"] == "succeeded"
    with Session(get_engine()) as db:
        assert db.scalar(select(func.count()).select_from(Document)) == 1


def test_docx_preserves_section_order_nested_tables_and_merged_cells(client: TestClient) -> None:
    doc = Docx()
    doc.add_paragraph("Synthetic teacher")
    doc.add_paragraph("Employment")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "School Alpha"
    table.cell(0, 1).text = "2019–2022"
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = "Curriculum lead"
    merged.add_table(rows=1, cols=1).cell(0, 0).text = "Mentored six teachers"
    doc.add_paragraph("Education")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic teaching qualification"
    doc.add_paragraph("Languages")
    buffer = io.BytesIO()
    doc.save(buffer)
    draft = draft_for(client, upload(client, auth(client), buffer.getvalue()))
    text = "\n".join(span["text"] for span in draft["spans"])
    ordered = [
        "Synthetic teacher",
        "Employment",
        "School Alpha",
        "2019–2022",
        "Curriculum lead",
        "Mentored six teachers",
        "Education",
        "Synthetic teaching qualification",
        "Languages",
    ]
    positions = [text.index(value) for value in ordered]
    assert positions == sorted(positions)
    assert text.count("Curriculum lead") == 1
    assert text.count("Mentored six teachers") == 1


def test_document_management_retention_and_cancellation(client: TestClient) -> None:
    assert client.get("/api/v1/profile/documents").status_code == 401
    headers = auth(client)
    pending = upload(client, headers)
    url = "/api/v1/profile/documents/" + pending["resource_id"]
    page = client.get("/api/v1/profile/documents").json()
    assert len(page["documents"]) == 1 and page["next_offset"] is None
    assert "private-name" not in str(page) and "Synthetic product designer" not in str(page)
    assert (
        client.patch(
            url + "/retention", headers=headers, json={"retain_original": True}
        ).status_code
        == 204
    )
    assert client.get("/api/v1/profile/documents").json()["documents"][0]["retain_original"] is True
    assert client.delete(url, headers=headers).status_code == 204
    process_operation(UUID(pending["operation_id"]))
    assert client.get(pending["status_url"]).json()["state"] == "cancelled"
    assert (
        client.patch(
            url + "/retention", headers=headers, json={"retain_original": True}
        ).status_code
        == 404
    )
    assert client.get(url + "/download").status_code == 404


def test_purge_recreate_does_not_reuse_revision(client: TestClient) -> None:
    headers = auth(client)
    body = {"facts": {"summary": "Synthetic nurse"}, "preferences": {"desired_roles": ["Nursing"]}}
    original = client.patch(
        "/api/v1/profile", headers={**headers, "If-Match": '"0"'}, json=body
    ).json()
    assert original["revision"] == 1
    assert (
        client.post(
            "/api/v1/profile/purge", headers=headers, json={"confirmation": "DELETE PROFILE"}
        ).status_code
        == 204
    )
    empty = client.get("/api/v1/profile").json()
    assert empty["revision"] > original["revision"] and empty["revision_id"] is None
    assert client.get("/api/v1/status").json()["profile_configured"] is False
    fresh = client.patch(
        "/api/v1/profile", headers={**headers, "If-Match": f'"{empty["revision"]}"'}, json=body
    )
    assert fresh.status_code == 200 and fresh.json()["revision"] > empty["revision"]
    assert client.get("/api/v1/status").json()["profile_configured"] is True
    assert (
        client.patch(
            "/api/v1/profile", headers={**headers, "If-Match": '"1"'}, json=body
        ).status_code
        == 409
    )
