"""Bounded document parser subprocess. No application settings or network clients."""

import json
import resource
import sys
import zipfile
from pathlib import Path

MAX_TEXT = 200_000
MAX_EXPANDED = 80 * 1024 * 1024


class DocumentError(ValueError):
    pass


def detect_type(data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return "application/pdf"
    if data.startswith(b"PK\x03\x04"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentError("Use a text PDF, DOCX, or UTF-8 text file.") from error
    if "\x00" in text or any(ord(char) < 32 and char not in "\t\r\n" for char in text):
        raise DocumentError("The file is not supported text.")
    return "text/plain"


def extract(path: Path, media_type: str) -> list[dict[str, str | int | None]]:
    segments: list[tuple[str, int | None, int | None]] = []
    if media_type == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentError("Upload an unencrypted PDF.")
        if len(reader.pages) > 50:
            raise DocumentError("The PDF exceeds the 50-page limit.")
        for index, page in enumerate(reader.pages):
            segments.append((page.extract_text() or "", index + 1, None))
    elif media_type.endswith("wordprocessingml.document"):
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 4000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED:
                raise DocumentError("Expanded DOCX content exceeds the safety limit.")
            names = {entry.filename for entry in entries}
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise DocumentError("The archive is not a DOCX document.")
            if any("vbaproject" in name.lower() for name in names):
                raise DocumentError("Macro-enabled documents are not supported.")
        from docx import Document

        doc = Document(str(path))
        for index, paragraph in enumerate(doc.paragraphs):
            segments.append((paragraph.text, None, index + 1))
        # Tables often contain the work history; do not silently omit them.
        for table in doc.tables:
            for row in table.rows:
                segments.append(
                    (" | ".join(cell.text for cell in row.cells), None, len(segments) + 1)
                )
    else:
        for index, text in enumerate(path.read_text(encoding="utf-8-sig").splitlines()):
            segments.append((text, None, index + 1))
    if sum(len(text) for text, _, _ in segments) > MAX_TEXT:
        raise DocumentError(
            "Extracted text is too large to process completely. Upload a shorter CV."
        )
    spans = [
        {"id": f"s{index + 1}", "text": text.strip(), "page": page, "paragraph": paragraph}
        for index, (text, page, paragraph) in enumerate(segments)
        if text.strip()
    ]
    if not spans:
        raise DocumentError(
            "No readable text was found. Upload a text-based CV; OCR is not supported."
        )
    return spans


def main() -> None:
    resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
    resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
    try:
        spans = extract(Path(sys.argv[1]), sys.argv[2])
        print(json.dumps({"spans": spans}))
    except DocumentError as error:
        print(json.dumps({"error": str(error)}))
        raise SystemExit(2) from None
    except Exception:
        # Parser tracebacks may contain private document text; never pass them into task logs.
        print(json.dumps({"error": "The document could not be read. Try a text PDF, DOCX or TXT."}))
        raise SystemExit(2) from None


if __name__ == "__main__":
    main()
